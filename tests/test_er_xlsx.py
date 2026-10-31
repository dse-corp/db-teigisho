from __future__ import annotations

import math
import posixpath
import warnings
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image

from db_teigisho.er_graph import build_er_graph
from db_teigisho.models import DatabaseDefinition
from db_teigisho.render import render_xlsx

NS = {
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
}


def _drawing(archive: ZipFile, title: str = "ER図") -> ET.Element:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    sheets = workbook.findall("s:sheets/s:sheet", NS)
    sheet = next(sheet for sheet in sheets if sheet.get("name") == title)
    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rel = next(rel for rel in rels if rel.get("Id") == sheet.get(f"{{{NS['r']}}}id"))
    target = rel.attrib["Target"]
    part = target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target)
    root = ET.fromstring(archive.read(part))
    reference = root.find("s:drawing", NS)
    assert reference is not None, "ER diagram must be a native worksheet drawing"
    relpath = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
    rels = ET.fromstring(archive.read(relpath))
    rel = next(rel for rel in rels if rel.get("Id") == reference.get(f"{{{NS['r']}}}id"))
    target = rel.attrib["Target"]
    drawing = (
        target.lstrip("/")
        if target.startswith("/")
        else posixpath.normpath(posixpath.join(posixpath.dirname(part), target))
    )
    content_types = ET.fromstring(archive.read("[Content_Types].xml"))
    namespaces = dict(
        item
        for _, item in ET.iterparse(
            BytesIO(archive.read("[Content_Types].xml")), events=["start-ns"]
        )
    )
    assert namespaces.get("") == NS["ct"], "OPC consumers require the default Types namespace"
    assert any(
        item.get("PartName") == "/" + drawing
        and item.get("ContentType") == "application/vnd.openxmlformats-officedocument.drawing+xml"
        for item in content_types
    )
    assert len(archive.namelist()) == len(set(archive.namelist()))
    return ET.fromstring(archive.read(drawing))


def _render(data: dict[str, Any], path: Path) -> ET.Element:
    render_xlsx(DatabaseDefinition.model_validate(data), path)
    with ZipFile(path) as archive:
        assert not any(name.startswith("xl/media/") for name in archive.namelist())
        return _drawing(archive)


def _assert_connections(root: ET.Element, count: int) -> None:
    ids = [item.attrib["id"] for item in root.findall(".//xdr:cNvPr", NS)]
    assert len(ids) == len(set(ids))
    shapes = {
        shape.find("xdr:nvSpPr/xdr:cNvPr", NS).attrib["id"]: shape  # type: ignore[union-attr]
        for shape in root.findall(".//xdr:sp", NS)
    }
    connectors = root.findall(".//xdr:cxnSp", NS)
    assert len(connectors) == count
    for connector in connectors:
        for tag in ("a:stCxn", "a:endCxn"):
            connection = connector.find(f"xdr:nvCxnSpPr/xdr:cNvCxnSpPr/{tag}", NS)
            assert connection is not None
            target = shapes[connection.attrib["id"]]
            sites = target.findall("xdr:spPr/a:custGeom/a:cxnLst/a:cxn", NS)
            preset = target.find("xdr:spPr/a:prstGeom", NS)
            site_count = 4 if preset is not None and preset.get("prst") == "rect" else len(sites)
            assert 0 <= int(connection.attrib["idx"]) < site_count
        assert (
            connector.find("xdr:spPr/a:prstGeom", NS) is not None
            or connector.find("xdr:spPr/a:custGeom/a:pathLst/a:path", NS) is not None
        )


def test_excel_er_contains_editable_groups_and_connected_elbows(
    valid_definition: dict[str, Any], tmp_path: Path
) -> None:
    root = _render(valid_definition, tmp_path / "er.xlsx")
    groups = root.findall("xdr:absoluteAnchor/xdr:grpSp", NS)
    assert len(groups) == 2
    assert not root.findall(".//xdr:pic", NS)
    _assert_connections(root, 1)
    for table, group in zip(valid_definition["tables"], groups, strict=True):
        texts = "".join(item.text or "" for item in group.findall(".//a:t", NS))
        assert table["physical_name"] in texts
        assert table["logical_name"] in texts
        for column in table["columns"]:
            assert column["physical_name"] in texts
            assert column["logical_name"] in texts
            assert column["data_type"] in texts
        assert "[1]" in texts
    assert root.find("xdr:absoluteAnchor/xdr:cxnSp/xdr:spPr/a:prstGeom", NS) is not None
    assert root.find(".//a:prstDash[@val='dash']", NS) is not None
    assert {"PK", "FK", "UK"} <= set(
        " ".join(item.text or "" for item in root.findall(".//a:t", NS)).split()
    )


def test_excel_er_does_not_render_png(
    definition_file: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from db_teigisho.cli import main

    def forbidden(*args: object, **kwargs: object) -> None:
        pytest.fail("XLSX generation must not invoke the image renderer")

    monkeypatch.setattr("db_teigisho.render.render_er_diagrams", forbidden)
    monkeypatch.setattr("db_teigisho.diagram_render._png", forbidden)
    assert (
        main(
            [
                "render",
                str(definition_file),
                "--format",
                "xlsx",
                "--output",
                str(tmp_path / "er.xlsx"),
            ]
        )
        == 0
    )


def test_excel_er_preserves_self_cyclic_parallel_and_composite_foreign_keys(
    valid_definition: dict[str, Any], tmp_path: Path
) -> None:
    parent, child = valid_definition["tables"]
    parent["foreign_keys"] = [
        {
            "name": "self",
            "columns": ["customer_id"],
            "referenced_table": "customers",
            "referenced_columns": ["customer_id"],
        },
        {
            "name": "cycle",
            "columns": ["customer_id"],
            "referenced_table": "orders",
            "referenced_columns": ["order_id"],
        },
    ]
    child["foreign_keys"].extend(
        [
            {**child["foreign_keys"][0], "name": "parallel"},
            {
                "name": "composite",
                "columns": ["customer_id", "order_id"],
                "referenced_table": "customers",
                "referenced_columns": ["customer_id", "email"],
                "on_delete": "CASCADE",
            },
        ]
    )
    root = _render(valid_definition, tmp_path / "complex.xlsx")
    _assert_connections(root, 5)
    connections = root.findall(".//xdr:cNvCxnSpPr", NS)
    endpoints = [
        tuple((child.get("id"), child.get("idx")) for child in item) for item in connections
    ]
    assert len(set(endpoints)) == 5
    assert root.find(".//a:prstDash[@val='solid']", NS) is not None
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="DrawingML support.*", category=UserWarning)
        workbook = load_workbook(tmp_path / "complex.xlsx")
    values = [str(cell.value) for row in workbook["ER図"] for cell in row if cell.value is not None]
    assert any("customer_id → customer_id" in value for value in values)
    assert any("order_id → email" in value for value in values)
    assert "composite" in values


@pytest.mark.parametrize("empty", [False, True])
def test_excel_er_supports_no_foreign_keys_and_no_tables(
    valid_definition: dict[str, Any], tmp_path: Path, empty: bool
) -> None:
    for table in valid_definition["tables"]:
        table["foreign_keys"] = []
    if empty:
        valid_definition["tables"] = []
    root = _render(valid_definition, tmp_path / "empty.xlsx")
    _assert_connections(root, 0)
    assert len(root.findall(".//xdr:grpSp", NS)) == (0 if empty else 2)
    if empty:
        assert "テーブル定義なし" in "".join(item.text or "" for item in root.findall(".//a:t", NS))


def test_excel_er_escapes_unicode_and_retains_long_names(
    valid_definition: dict[str, Any], tmp_path: Path
) -> None:
    name = '日本語<&>" の長い名前' * 25
    valid_definition["tables"][0]["logical_name"] = name
    root = _render(valid_definition, tmp_path / "unicode.xlsx")
    texts = "".join(item.text or "" for item in root.findall(".//a:t", NS))
    assert name in texts
    for extent in root.findall(".//a:ext", NS) + root.findall(".//xdr:ext", NS):
        assert int(extent.attrib["cx"]) >= 0
        assert int(extent.attrib["cy"]) >= 0


def test_excel_er_print_area_contains_diagram_and_legend(
    valid_definition: dict[str, Any], tmp_path: Path
) -> None:
    for i in range(10):
        table = deepcopy(valid_definition["tables"][0])
        table["physical_name"] = f"isolated_{i}"
        valid_definition["tables"].append(table)
    root = _render(valid_definition, tmp_path / "large.xlsx")
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="DrawingML support.*", category=UserWarning)
        workbook = load_workbook(tmp_path / "large.xlsx")
    sheet = workbook["ER図"]
    from openpyxl.utils.cell import range_boundaries

    area = sheet.print_area.rsplit("!", 1)[-1]
    _, _, end_col, end_row = range_boundaries(area)
    assert end_row > 45
    assert end_col >= 16
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.page_setup.fitToHeight == 0
    assert len(root.findall(".//xdr:grpSp", NS)) == 12


def test_drawing_injection_preserves_other_drawings_and_allocates_names(
    valid_definition: dict[str, Any], tmp_path: Path
) -> None:
    from db_teigisho.er_xlsx import build_er_drawing, save_with_er_drawing

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Other"
    buffer = BytesIO()
    Image.new("RGB", (4, 4), "red").save(buffer, format="PNG")
    buffer.seek(0)
    sheet.add_image(ExcelImage(buffer), "B2")
    er_sheet = workbook.create_sheet("ER図")
    er_sheet["A1"] = "keep me"
    er_sheet["B1"].hyperlink = "https://example.com"
    output = tmp_path / "injected.xlsx"
    drawing = build_er_drawing(build_er_graph(DatabaseDefinition.model_validate(valid_definition)))
    save_with_er_drawing(workbook, er_sheet.title, drawing, output)
    with ZipFile(output) as archive:
        assert "xl/drawings/drawing1.xml" in archive.namelist()
        assert "xl/drawings/drawing2.xml" in archive.namelist()
        _assert_connections(_drawing(archive), 1)
        rels = ET.fromstring(archive.read("xl/worksheets/_rels/sheet2.xml.rels"))
        assert len(rels) == 2
        assert len({rel.get("Id") for rel in rels}) == 2
        assert any(rel.get("Target") == "https://example.com" for rel in rels)
        assert (
            ET.fromstring(archive.read("xl/drawings/drawing1.xml")).find(".//xdr:pic", NS)
            is not None
        )


def test_failed_drawing_save_keeps_previous_output_and_removes_tempfile(
    valid_definition: dict[str, Any], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from db_teigisho.er_xlsx import build_er_drawing, save_with_er_drawing

    workbook = Workbook()
    assert workbook.active is not None
    workbook.active.title = "ER図"
    drawing = build_er_drawing(build_er_graph(DatabaseDefinition.model_validate(valid_definition)))
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"previous output")

    def fail(*args: object, **kwargs: object) -> None:
        raise OSError("disk full")

    monkeypatch.setattr("db_teigisho.er_xlsx.os.replace", fail)
    with pytest.raises(OSError, match="disk full"):
        save_with_er_drawing(workbook, "ER図", drawing, output)
    assert output.read_bytes() == b"previous output"
    assert list(tmp_path.iterdir()) == [output]


def test_very_long_foreign_key_legend_continues_without_clipping(
    valid_definition: dict[str, Any], tmp_path: Path
) -> None:
    name = "非常に長い外部キー名<&>" * 1000
    valid_definition["tables"][1]["foreign_keys"][0]["name"] = name
    output = tmp_path / "long-legend.xlsx"
    _render(valid_definition, output)
    workbook = load_workbook(output)
    sheet = workbook["ER図"]
    parts = [row[1].value for row in sheet.iter_rows() if row[1].value is not None]
    assert "".join(parts).replace("\n", "").endswith(name)
    from db_teigisho.er_scene import text_width

    for row in sheet.iter_rows():
        if row[1].value is None:
            continue
        value = str(row[1].value)
        required_lines = sum(
            max(1, math.ceil(text_width(line, 14) / 253)) for line in value.splitlines()
        )
        assert required_lines * 16 <= sheet.row_dimensions[row[1].row].height
    workbook.close()


@pytest.mark.parametrize(
    ("nullable", "unique", "cascade"),
    [
        (False, False, False),
        (True, False, False),
        (False, True, True),
    ],
)
def test_endpoint_cardinality_and_line_style_match_graph(
    valid_definition: dict[str, Any], tmp_path: Path, nullable: bool, unique: bool, cascade: bool
) -> None:
    child = valid_definition["tables"][1]
    child["columns"][1]["not_null"] = not nullable
    child["columns"][1]["unique"] = unique
    child["foreign_keys"][0]["on_delete"] = "CASCADE" if cascade else "NO ACTION"
    root = _render(valid_definition, tmp_path / "cardinality.xlsx")
    endpoint_values = {
        item.text
        for shape in root.findall(".//xdr:sp", NS)
        if shape.find("xdr:nvSpPr/xdr:cNvPr", NS).get("name", "").startswith("Endpoint:")  # type: ignore[union-attr]
        for item in shape.findall(".//a:t", NS)
    }
    assert endpoint_values == {
        f"[1] {'0..1' if nullable else '1'}",
        f"[1] {'0..1' if unique else '0..N'}",
    }
    assert (
        root.find(
            f".//xdr:cxnSp/xdr:spPr/a:ln/a:prstDash[@val='{'solid' if cascade else 'dash'}']", NS
        )
        is not None
    )


def test_excel_er_rejects_limits_before_saving_output(
    valid_definition: dict[str, Any], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from db_teigisho.errors import DefinitionRenderError

    monkeypatch.setattr("db_teigisho.er_scene.MAX_TABLES", 1)
    output = tmp_path / "oversized.xlsx"
    with pytest.raises(DefinitionRenderError, match="safety limits"):
        render_xlsx(DatabaseDefinition.model_validate(valid_definition), output)
    assert not output.exists()


def test_excel_er_handles_2000_distinct_parallel_foreign_keys(
    valid_definition: dict[str, Any],
) -> None:
    from db_teigisho.er_xlsx import build_er_drawing

    child = valid_definition["tables"][1]
    foreign_key = child["foreign_keys"][0]
    child["foreign_keys"] = [{**foreign_key, "name": f"fk_{i}"} for i in range(2000)]
    drawing = build_er_drawing(build_er_graph(DatabaseDefinition.model_validate(valid_definition)))
    root = ET.fromstring(drawing.xml)
    _assert_connections(root, 2000)
    sites = root.findall(".//xdr:cNvCxnSpPr/a:stCxn", NS)
    assert len({site.attrib["id"] for site in sites}) == 2000
    assert len(drawing.legend) == 2000
    assert drawing.height < 100000


def test_excel_er_handles_500_tables_in_a_cycle(valid_definition: dict[str, Any]) -> None:
    from db_teigisho.er_xlsx import build_er_drawing

    template = valid_definition["tables"][0]
    template["columns"] = template["columns"][:1]
    template["indexes"] = []
    tables = []
    for i in range(500):
        table = deepcopy(template)
        table["physical_name"] = f"t_{i:03}"
        table["foreign_keys"] = [
            {
                "name": f"fk_{i}",
                "columns": ["customer_id"],
                "referenced_table": f"t_{(i + 1) % 500:03}",
                "referenced_columns": ["customer_id"],
            }
        ]
        tables.append(table)
    valid_definition["tables"] = tables
    drawing = build_er_drawing(build_er_graph(DatabaseDefinition.model_validate(valid_definition)))
    root = ET.fromstring(drawing.xml)
    assert len(root.findall(".//xdr:grpSp", NS)) == 500
    _assert_connections(root, 500)
    assert drawing.height < 100000


def _world_points(properties: ET.Element) -> list[tuple[float, float]]:
    """Evaluate the written DrawingML geometry, independently of the router."""
    transform = properties.find("a:xfrm", NS)
    assert transform is not None
    off = transform.find("a:off", NS)
    ext = transform.find("a:ext", NS)
    assert off is not None and ext is not None
    x, y = float(off.attrib["x"]), float(off.attrib["y"])
    w, h = float(ext.attrib["cx"]), float(ext.attrib["cy"])
    preset = properties.find("a:prstGeom", NS)
    if preset is not None:
        values = [
            int(gd.attrib["fmla"].split()[1]) / 100000 for gd in preset.findall("a:avLst/a:gd", NS)
        ]
        x1 = w * values[0]
        kind = preset.attrib["prst"]
        if kind == "bentConnector3":
            points = [(0, 0), (x1, 0), (x1, h), (w, h)]
        elif kind == "bentConnector4":
            y2 = h * values[1]
            points = [(0, 0), (x1, 0), (x1, y2), (w, y2), (w, h)]
        else:
            y2, x3 = h * values[1], w * values[2]
            points = [(0, 0), (x1, 0), (x1, y2), (x3, y2), (x3, h), (w, h)]
    else:
        geometry = properties.find("a:custGeom", NS)
        assert geometry is not None
        guides = {"w": w, "h": h}
        for guide in geometry.findall("a:avLst/a:gd", NS) + geometry.findall("a:gdLst/a:gd", NS):
            operation, *args = guide.attrib["fmla"].split()
            values_ = [guides[value] if value in guides else float(value) for value in args]
            guides[guide.attrib["name"]] = (
                values_[0] if operation == "val" else values_[0] * values_[1] / values_[2]
            )
        points = [
            (guides[pt.attrib["x"]], guides[pt.attrib["y"]])
            for pt in geometry.findall("a:pathLst/a:path/*/a:pt", NS)
        ]
    result = []
    angle = math.radians(int(transform.get("rot", "0")) / 60000)
    for px, py in points:
        px = w - px if transform.get("flipH") == "1" else px
        py = h - py if transform.get("flipV") == "1" else py
        rx = (px - w / 2) * math.cos(angle) - (py - h / 2) * math.sin(angle)
        ry = (px - w / 2) * math.sin(angle) + (py - h / 2) * math.cos(angle)
        point = ((x + w / 2 + rx) / 9525, (y + h / 2 + ry) / 9525)
        if not result or math.dist(result[-1], point) > 0.01:
            result.append(point)
    return result


@pytest.mark.parametrize(
    "route",
    [
        ((10, 20), (70, 20)),
        ((10, 20), (70, 20), (70, 90)),
        ((10, 20), (40, 20), (40, 90), (70, 90)),
        ((10, 20), (40, 20), (40, 140), (70, 140), (70, 90)),
        ((10, 20), (40, 20), (40, 140), (60, 140), (60, 90), (70, 90)),
        ((10, 20), (40, 20), (40, 90), (10, 90)),
        ((10, 20), (40, 20), (40, 30), (50, 30), (50, 60), (60, 60), (60, 90), (70, 90)),
    ],
)
@pytest.mark.parametrize(
    "mirror_x,mirror_y,vertical",
    [
        (False, False, False),
        (True, True, False),
        (True, False, False),
        (False, True, False),
        (False, False, True),
        (True, True, True),
    ],
)
def test_serialized_connector_preserves_route_in_every_direction(
    route: tuple[tuple[float, float], ...], mirror_x: bool, mirror_y: bool, vertical: bool
) -> None:
    from db_teigisho.er_scene import Edge
    from db_teigisho.er_xlsx import _connector, _Drawing

    points = tuple((200 - x if mirror_x else x, 200 - y if mirror_y else y) for x, y in route)
    if vertical:
        points = tuple((y, x) for x, y in points)
    root = ET.Element("drawing")
    _connector(_Drawing(), root, Edge("test", points, False), "test", (1, 0), (2, 0), 0, 0)
    properties = root.find(".//xdr:cxnSp/xdr:spPr", NS)
    assert properties is not None
    actual = _world_points(properties)
    if len(points) == 2:
        # A collapsed elbow has a collinear adjustment point on the same line.
        assert sum(
            math.dist(a, b) for a, b in zip(actual, actual[1:], strict=False)
        ) == pytest.approx(math.dist(*points))
        actual = [actual[0], actual[-1]]
    assert len(actual) == len(points)
    for expected, got in zip(points, actual, strict=True):
        assert math.dist(expected, got) < 0.01


def test_custom_bend_handles_preserve_orthogonality_and_endpoints() -> None:
    from db_teigisho.er_xlsx import _custom_connector

    properties = ET.Element("properties")
    _custom_connector(properties, ((10, 20), (40, 20), (40, 90), (10, 90)))
    before = _world_points(properties)
    handle = properties.find("a:custGeom/a:ahLst/a:ahXY", NS)
    assert handle is not None
    adjustment = handle.attrib["gdRefX"]
    guide = properties.find(f"a:custGeom/a:avLst/a:gd[@name='{adjustment}']", NS)
    assert guide is not None
    guide.set("fmla", "val 800000")
    after = _world_points(properties)
    assert before != after
    assert before[0] == after[0] and before[-1] == after[-1]
    assert all(a[0] == b[0] or a[1] == b[1] for a, b in zip(after, after[1:], strict=False))


def test_connector_geometry_starts_and_ends_at_the_attached_table_ports(
    valid_definition: dict[str, Any], tmp_path: Path
) -> None:
    root = _render(valid_definition, tmp_path / "ports.xlsx")
    ports = {}
    for group in root.findall(".//xdr:grpSp", NS):
        transform = group.find("xdr:grpSpPr/a:xfrm", NS)
        assert transform is not None
        off = transform.find("a:off", NS)
        child_off = transform.find("a:chOff", NS)
        assert off is not None and child_off is not None
        gx = int(off.attrib["x"]) - int(child_off.attrib["x"])
        gy = int(off.attrib["y"]) - int(child_off.attrib["y"])
        for shape in group.findall("xdr:sp", NS):
            identity = shape.find("xdr:nvSpPr/xdr:cNvPr", NS)
            assert identity is not None
            if not identity.attrib["name"].startswith("Port:"):
                continue
            offset = shape.find("xdr:spPr/a:xfrm/a:off", NS)
            extent = shape.find("xdr:spPr/a:xfrm/a:ext", NS)
            assert offset is not None and extent is not None
            x, y = gx + int(offset.attrib["x"]), gy + int(offset.attrib["y"])
            w, h = int(extent.attrib["cx"]), int(extent.attrib["cy"])
            ports[identity.attrib["id"]] = [
                (x + w / 2, y),
                (x, y + h / 2),
                (x + w / 2, y + h),
                (x + w, y + h / 2),
            ]
    for connector in root.findall(".//xdr:cxnSp", NS):
        properties = connector.find("xdr:spPr", NS)
        assert properties is not None
        path = _world_points(properties)
        for tag, point in (("stCxn", path[0]), ("endCxn", path[-1])):
            connection = connector.find(f"xdr:nvCxnSpPr/xdr:cNvCxnSpPr/a:{tag}", NS)
            assert connection is not None
            target = ports[connection.attrib["id"]][int(connection.attrib["idx"])]
            assert math.dist(point, (target[0] / 9525, target[1] / 9525)) < 0.01


def test_parallel_fk_labels_on_top_and_bottom_do_not_overlap(
    valid_definition: dict[str, Any], tmp_path: Path
) -> None:
    parent, child = valid_definition["tables"]
    parent["foreign_keys"] = [
        {
            "name": "cycle",
            "columns": ["customer_id"],
            "referenced_table": "orders",
            "referenced_columns": ["order_id"],
        }
    ]
    original = child["foreign_keys"][0]
    child["foreign_keys"] = [{**original, "name": f"parallel_{i}"} for i in range(25)]
    root = _render(valid_definition, tmp_path / "dense-cycle.xlsx")
    for group in root.findall(".//xdr:grpSp", NS):
        boxes = []
        for shape in group.findall("xdr:sp", NS):
            identity = shape.find("xdr:nvSpPr/xdr:cNvPr", NS)
            assert identity is not None
            if not identity.get("name", "").startswith("Endpoint:"):
                continue
            offset = shape.find("xdr:spPr/a:xfrm/a:off", NS)
            extent = shape.find("xdr:spPr/a:xfrm/a:ext", NS)
            assert offset is not None and extent is not None
            x, y = int(offset.attrib["x"]), int(offset.attrib["y"])
            w, h = int(extent.attrib["cx"]), int(extent.attrib["cy"])
            assert all(
                x + w <= bx or bx + bw <= x or y + h <= by or by + bh <= y
                for bx, by, bw, bh in boxes
            )
            boxes.append((x, y, w, h))
